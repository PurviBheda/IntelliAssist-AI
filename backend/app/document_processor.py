"""IntelliAssist AI - Document Processing Module.

This module handles extraction of structured text from PDF, TXT, and DOCX documents.
It preserves PDF page numbers, represents TXT and DOCX as single logical pages,
calculates character and page counts, and provides robust error handling.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Dict, List, Union
import zipfile

import pymupdf as fitz  # PyMuPDF
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


def _process_txt(path: Path) -> List[Dict[str, Any]]:
    """Extract text from a UTF-8 encoded text file.

    Handles UTF-8 Byte Order Marks (BOM) automatically.
    Surrounding whitespace is stripped.
    Represents the content as one logical page (page_number 1).

    Raises:
        ValueError: If file is empty, contains only whitespace, or cannot be decoded as UTF-8.
    """
    try:
        raw_text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as e:
        raise ValueError(
            f"Unable to decode text file '{path.name}' as UTF-8: {e}"
        ) from e

    cleaned_text = raw_text.strip()
    if not cleaned_text:
        raise ValueError(
            f"TXT document is empty or contains only whitespace: {path.name}"
        )

    return [{"page_number": 1, "text": cleaned_text}]


def _process_docx(path: Path) -> List[Dict[str, Any]]:
    """Extract text from a Word DOCX file.

    Extracts non-empty paragraphs and table cell contents in reading order.
    Represents the content as one logical page (page_number 1).

    Raises:
        ValueError: If file is malformed, corrupted, or contains no readable text.
    """
    try:
        doc = Document(str(path))
    except (zipfile.BadZipFile, KeyError, Exception) as e:
        raise ValueError(
            f"Invalid or corrupted DOCX document '{path.name}': {e}"
        ) from e

    content_blocks: List[str] = []

    # Iterate through body elements to preserve inter-paragraph and table reading order
    try:
        for child in doc.element.body:
            if child.tag.endswith("p"):
                paragraph = Paragraph(child, doc)
                text = paragraph.text.strip()
                if text:
                    content_blocks.append(text)
            elif child.tag.endswith("tbl"):
                table = Table(child, doc)
                for row in table.rows:
                    seen_cells = set()
                    row_cells_text: List[str] = []
                    for cell in row.cells:
                        if cell._tc not in seen_cells:
                            seen_cells.add(cell._tc)
                            c_text = cell.text.strip()
                            if c_text:
                                row_cells_text.append(c_text)
                    if row_cells_text:
                        content_blocks.append(" | ".join(row_cells_text))
    except Exception:
        # Fallback to direct paragraph and table traversal if element iteration fails
        content_blocks.clear()
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                content_blocks.append(text)
        for tbl in doc.tables:
            for row in tbl.rows:
                seen_cells = set()
                row_cells_text = []
                for cell in row.cells:
                    if cell._tc not in seen_cells:
                        seen_cells.add(cell._tc)
                        c_text = cell.text.strip()
                        if c_text:
                            row_cells_text.append(c_text)
                if row_cells_text:
                    content_blocks.append(" | ".join(row_cells_text))

    full_text = "\n\n".join(content_blocks).strip()
    if not full_text:
        raise ValueError(
            f"DOCX document contains no readable text: {path.name}"
        )

    return [{"page_number": 1, "text": full_text}]


def _process_pdf(path: Path) -> List[Dict[str, Any]]:
    """Extract text from a PDF document using PyMuPDF.

    Iterates through all pages, preserving 1-based page numbers and page order.
    Resources are safely closed via context management.

    Raises:
        ValueError: If PDF is encrypted, corrupted, has no pages, or contains
                    no extractable text (e.g., scanned images requiring OCR).
    """
    try:
        doc = fitz.open(str(path))
    except Exception as e:
        raise ValueError(
            f"Invalid or unreadable PDF document '{path.name}': {e}"
        ) from e

    pages: List[Dict[str, Any]] = []

    with doc:
        if doc.is_encrypted and doc.needs_pass:
            raise ValueError(
                f"Password-protected PDF documents are not supported: {path.name}"
            )

        if doc.page_count == 0:
            raise ValueError(f"PDF document contains no pages: {path.name}")

        for page_index in range(doc.page_count):
            try:
                page = doc.load_page(page_index)
                page_text = page.get_text() or ""
                # Strip leading and trailing whitespace on the page text
                pages.append({
                    "page_number": page_index + 1,
                    "text": page_text.strip(),
                })
            except Exception as e:
                raise ValueError(
                    f"Failed to read page {page_index + 1} of PDF '{path.name}': {e}"
                ) from e

    # Check whether the PDF contained any extractable text across all pages
    total_text = "".join(p["text"] for p in pages).strip()
    if not total_text:
        raise ValueError(
            f"PDF document contains no extractable readable text: {path.name}. "
            "The document may contain only scanned images or raster graphics, "
            "which require OCR (Optical Character Recognition). "
            "OCR is not yet implemented in this milestone."
        )

    return pages


def process_document(file_path: Union[str, os.PathLike]) -> Dict[str, Any]:
    """Process a document file and return structured text and metadata.

    Supported formats:
        - PDF (.pdf): Extracts page-by-page text preserving 1-based page numbering.
        - TXT (.txt): Reads UTF-8 / UTF-8-BOM plain text as a single logical page.
        - DOCX (.docx): Extracts paragraphs and tables as a single logical page.

    Args:
        file_path: Filesystem path to the target document.

    Returns:
        A dictionary with the following schema:
        {
            "filename": str,
            "file_type": str,
            "pages": [
                {
                    "page_number": int,
                    "text": str
                }
            ],
            "total_pages": int,
            "total_characters": int
        }

    Raises:
        FileNotFoundError: If the file does not exist.
        ValueError: If the path is not a regular file, the file extension is
                    unsupported, the document is empty, or no readable text can be extracted.
    """
    if file_path is None:
        raise ValueError("file_path cannot be None")

    path = Path(file_path)

    # 1. Existence and regular file checks
    if not path.exists():
        raise FileNotFoundError(f"Document file not found: {file_path}")

    if not path.is_file():
        raise ValueError(f"Provided path is not a regular file: {file_path}")

    # 2. Extension validation (case-insensitive)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format: '{path.suffix}'. "
            f"Supported formats are: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    file_type = suffix.lstrip(".")

    # 3. Format-specific processing
    if suffix == ".txt":
        pages = _process_txt(path)
    elif suffix == ".docx":
        pages = _process_docx(path)
    elif suffix == ".pdf":
        pages = _process_pdf(path)
    else:  # Defensive safeguard
        raise ValueError(f"Unsupported file format: {suffix}")

    total_characters = sum(len(page["text"]) for page in pages)

    return {
        "filename": path.name,
        "file_type": file_type,
        "pages": pages,
        "total_pages": len(pages),
        "total_characters": total_characters,
    }

"""Automated unit tests for IntelliAssist AI Document Processor.

Covers PDF, TXT, and DOCX extraction, error handling, metadata validation,
and edge cases using pytest temporary directories and programmatic fixtures.
"""

from __future__ import annotations

import codecs
from pathlib import Path
from typing import List, Optional

import fitz  # PyMuPDF
import pytest
from docx import Document

from app.document_processor import process_document


# ============================================================================
# Fixture Helpers
# ============================================================================


def create_sample_pdf(path: Path, pages: List[str]) -> Path:
    """Programmatically generate a multi-page PDF document."""
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text(fitz.Point(72, 72), text)
    doc.save(str(path))
    doc.close()
    return path


def create_sample_docx(
    path: Path,
    paragraphs: List[str],
    table_rows: Optional[List[List[str]]] = None,
) -> Path:
    """Programmatically generate a Word DOCX document."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_idx, row in enumerate(table_rows):
            for c_idx, cell_value in enumerate(row):
                table.cell(r_idx, c_idx).text = cell_value
    doc.save(str(path))
    return path


# ============================================================================
# TXT Processing Tests (Scenarios 1-5, 8, and BOM)
# ============================================================================


def test_txt_successful_extraction(tmp_path: Path) -> None:
    """Scenario 1: Successful TXT extraction."""
    txt_file = tmp_path / "notes.txt"
    content = "Hello, IntelliAssist AI!"
    txt_file.write_text(content, encoding="utf-8")

    result = process_document(txt_file)

    assert isinstance(result, dict)
    assert result["filename"] == "notes.txt"
    assert result["file_type"] == "txt"
    assert len(result["pages"]) == 1
    assert result["pages"][0]["text"] == content


def test_txt_filename_and_file_type(tmp_path: Path) -> None:
    """Scenario 2: Correct filename and file type attributes."""
    txt_file = tmp_path / "custom_report_2026.txt"
    txt_file.write_text("Financial overview data", encoding="utf-8")

    result = process_document(txt_file)

    assert result["filename"] == "custom_report_2026.txt"
    assert result["file_type"] == "txt"


def test_txt_page_count(tmp_path: Path) -> None:
    """Scenario 3: Correct page count for TXT (logical page 1)."""
    txt_file = tmp_path / "multi_line.txt"
    txt_file.write_text("Line 1\nLine 2\nLine 3\nLine 4", encoding="utf-8")

    result = process_document(txt_file)

    assert result["total_pages"] == 1
    assert len(result["pages"]) == 1
    assert result["pages"][0]["page_number"] == 1


def test_txt_extracted_text(tmp_path: Path) -> None:
    """Scenario 4: Correct extracted text with whitespace stripping."""
    txt_file = tmp_path / "whitespace.txt"
    raw = "\n\n   Leading and trailing whitespace to strip   \n\n"
    txt_file.write_text(raw, encoding="utf-8")

    result = process_document(txt_file)

    assert result["pages"][0]["text"] == "Leading and trailing whitespace to strip"


def test_txt_character_count(tmp_path: Path) -> None:
    """Scenario 5: Correct character count matching extracted text length."""
    txt_file = tmp_path / "chars.txt"
    text = "Exact 25 characters here."
    assert len(text) == 25
    txt_file.write_text(text, encoding="utf-8")

    result = process_document(txt_file)

    assert result["total_characters"] == 25


def test_empty_txt_document(tmp_path: Path) -> None:
    """Scenario 8: Empty TXT document or whitespace-only raises ValueError."""
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("", encoding="utf-8")

    with pytest.raises(ValueError, match="empty or contains only whitespace"):
        process_document(empty_file)

    whitespace_file = tmp_path / "whitespace_only.txt"
    whitespace_file.write_text("   \n\t  \n  ", encoding="utf-8")

    with pytest.raises(ValueError, match="empty or contains only whitespace"):
        process_document(whitespace_file)


def test_txt_utf8_bom_handling(tmp_path: Path) -> None:
    """Handles UTF-8 Byte Order Mark (BOM) transparently."""
    bom_file = tmp_path / "bom.txt"
    content = "Document starting with UTF-8 BOM"
    bom_file.write_bytes(codecs.BOM_UTF8 + content.encode("utf-8"))

    result = process_document(bom_file)

    assert result["pages"][0]["text"] == content
    assert not result["pages"][0]["text"].startswith("\ufeff")


# ============================================================================
# Validation and Error Handling Tests (Scenarios 6, 7)
# ============================================================================


def test_unsupported_file_extension(tmp_path: Path) -> None:
    """Scenario 6: Unsupported file extension raises ValueError."""
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("col1,col2\n1,2", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file format"):
        process_document(csv_file)

    png_file = tmp_path / "image.png"
    png_file.write_bytes(b"\x89PNG\r\n\x1a\n")

    with pytest.raises(ValueError, match="Unsupported file format"):
        process_document(png_file)


def test_missing_file(tmp_path: Path) -> None:
    """Scenario 7: Missing file raises FileNotFoundError."""
    missing = tmp_path / "non_existent_document.pdf"

    with pytest.raises(FileNotFoundError, match="Document file not found"):
        process_document(missing)


def test_path_is_directory(tmp_path: Path) -> None:
    """Passing a directory instead of a regular file raises ValueError."""
    dir_path = tmp_path / "fake_dir.pdf"
    dir_path.mkdir()

    with pytest.raises(ValueError, match="not a regular file"):
        process_document(dir_path)


# ============================================================================
# DOCX Processing Tests (Scenarios 9, 10)
# ============================================================================


def test_docx_successful_extraction(tmp_path: Path) -> None:
    """Scenario 9: Successful DOCX extraction with paragraphs."""
    docx_file = tmp_path / "sample.docx"
    paragraphs = [
        "First paragraph of the project proposal.",
        "Second paragraph with additional technical specs.",
    ]
    create_sample_docx(docx_file, paragraphs)

    result = process_document(docx_file)

    assert result["filename"] == "sample.docx"
    assert result["file_type"] == "docx"
    assert result["total_pages"] == 1
    assert result["pages"][0]["page_number"] == 1
    assert "First paragraph of the project proposal." in result["pages"][0]["text"]
    assert "Second paragraph with additional technical specs." in result["pages"][0]["text"]
    assert result["total_characters"] == len(result["pages"][0]["text"])


def test_docx_table_text_extraction(tmp_path: Path) -> None:
    """Scenario 10: DOCX table text extraction with sensible reading order."""
    docx_file = tmp_path / "table_doc.docx"
    paragraphs = ["Executive Summary"]
    table_data = [
        ["Quarter", "Revenue", "Margin"],
        ["Q1", "$1.2M", "24%"],
        ["Q2", "$1.5M", "27%"],
    ]
    create_sample_docx(docx_file, paragraphs, table_rows=table_data)

    result = process_document(docx_file)

    extracted_text = result["pages"][0]["text"]
    assert "Executive Summary" in extracted_text
    assert "Quarter | Revenue | Margin" in extracted_text
    assert "Q1 | $1.2M | 24%" in extracted_text
    assert "Q2 | $1.5M | 27%" in extracted_text


def test_docx_empty_document(tmp_path: Path) -> None:
    """DOCX with no readable text raises ValueError."""
    empty_docx = tmp_path / "empty.docx"
    create_sample_docx(empty_docx, ["   ", ""])

    with pytest.raises(ValueError, match="contains no readable text"):
        process_document(empty_docx)


# ============================================================================
# PDF Processing Tests (Scenarios 11-15)
# ============================================================================


def test_pdf_successful_extraction(tmp_path: Path) -> None:
    """Scenario 11: Successful PDF extraction."""
    pdf_file = tmp_path / "document.pdf"
    create_sample_pdf(pdf_file, ["IntelliAssist AI Document Content"])

    result = process_document(pdf_file)

    assert result["filename"] == "document.pdf"
    assert result["file_type"] == "pdf"
    assert result["total_pages"] == 1
    assert result["pages"][0]["page_number"] == 1
    assert "IntelliAssist AI Document Content" in result["pages"][0]["text"]
    assert result["total_characters"] == len(result["pages"][0]["text"])


def test_pdf_page_numbering(tmp_path: Path) -> None:
    """Scenario 12: Correct 1-based PDF page numbering across multiple pages."""
    pdf_file = tmp_path / "three_pages.pdf"
    pages = ["Page 1 Intro", "Page 2 Methods", "Page 3 Conclusion"]
    create_sample_pdf(pdf_file, pages)

    result = process_document(pdf_file)

    assert result["total_pages"] == 3
    assert len(result["pages"]) == 3
    for idx, page_entry in enumerate(result["pages"], start=1):
        assert page_entry["page_number"] == idx


def test_pdf_page_order(tmp_path: Path) -> None:
    """Scenario 13: PDF page order is preserved sequentially."""
    pdf_file = tmp_path / "ordered_pages.pdf"
    pages = [
        "Alpha section content",
        "Beta section content",
        "Gamma section content",
    ]
    create_sample_pdf(pdf_file, pages)

    result = process_document(pdf_file)

    assert result["pages"][0]["text"] == "Alpha section content"
    assert result["pages"][1]["text"] == "Beta section content"
    assert result["pages"][2]["text"] == "Gamma section content"


def test_pdf_no_extractable_text(tmp_path: Path) -> None:
    """Scenario 14: PDF with no extractable text raises ValueError explaining OCR."""
    pdf_file = tmp_path / "scanned_or_blank.pdf"
    # Create PDF with 2 blank pages (no text inserted)
    create_sample_pdf(pdf_file, ["", ""])

    with pytest.raises(ValueError, match="OCR"):
        process_document(pdf_file)


def test_pdf_invalid_or_damaged(tmp_path: Path) -> None:
    """Scenario 15: Invalid or damaged PDF raises a descriptive ValueError."""
    corrupted_pdf = tmp_path / "corrupted.pdf"
    corrupted_pdf.write_bytes(b"%PDF-1.4\nCorrupted binary payload\x00\xff")

    with pytest.raises(ValueError, match="Invalid or unreadable PDF document"):
        process_document(corrupted_pdf)


def test_case_insensitive_extension(tmp_path: Path) -> None:
    """File extension matching is case-insensitive (.PDF, .TXT, .DOCX)."""
    uppercase_pdf = tmp_path / "uppercase.PDF"
    create_sample_pdf(uppercase_pdf, ["Uppercase extension test"])

    result = process_document(uppercase_pdf)
    assert result["file_type"] == "pdf"
    assert result["pages"][0]["text"] == "Uppercase extension test"
